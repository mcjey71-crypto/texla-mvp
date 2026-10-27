"""Developer-only fixture generation (python-docx, reportlab and Pillow)."""
from pathlib import Path
from io import BytesIO
from docx import Document
from docx.oxml import parse_xml
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from reportlab.lib.pdfencrypt import StandardEncryption
from PIL import Image, ImageDraw, ImageFont
from zipfile import ZipFile, ZIP_DEFLATED
root = Path('tests/fixtures/import')
root.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_heading('Investigación de prueba', 1)
p = doc.add_paragraph('Un resultado ')
p.add_run('importante').bold = True
p.add_run(' para la tesis.')
doc.add_paragraph('Primera observación', style='List Bullet')
doc.add_paragraph('Segunda observación', style='List Bullet')
table = doc.add_table(rows=2, cols=2)
for row, values in zip(table.rows, [('Variable','Valor'),('Masa','12')]):
    for cell, value in zip(row.cells, values): cell.text = value
p = doc.add_paragraph()
p._p.append(parse_xml('<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f></m:oMath></m:oMathPara>'))
doc.save(root / 'research.docx')
canvas = Canvas(str(root/'research.pdf'))
canvas.setFont('Helvetica-Bold',18);canvas.drawString(60,780,'Investigacion de prueba')
canvas.setFont('Helvetica',12);canvas.drawString(60,740,'Este documento contiene texto seleccionable.')
canvas.drawString(60,724,'La segunda linea forma parte del mismo parrafo.')
canvas.showPage();canvas.setFont('Helvetica',12);canvas.drawString(60,740,'Resultados de la segunda pagina.');canvas.save()
image = Image.new('RGB',(1400,350),'white');draw = ImageDraw.Draw(image)
font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',45)
draw.text((60,70),'Documento escaneado de prueba',font=font,fill='black')
draw.text((60,145),'La ciencia necesita evidencia.',font=font,fill='black')
canvas=Canvas(str(root/'scanned.pdf'));canvas.drawImage(ImageReader(image),40,530,width=520,height=130);canvas.save()
canvas=Canvas(str(root/'protected.pdf'),encrypt=StandardEncryption('test-only-password'));canvas.drawString(50,750,'Private fixture');canvas.save()
canvas=Canvas(str(root/'too-many-pages.pdf'))
for i in range(101): canvas.drawString(50,750,str(i));canvas.showPage()
canvas.save()
with ZipFile(root/'research.odt','w',ZIP_DEFLATED) as archive:
    archive.writestr('mimetype','application/vnd.oasis.opendocument.text')
    archive.writestr('content.xml','''<?xml version="1.0"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text><text:h text:outline-level="1">Informe abierto</text:h><text:p>Contenido de OpenDocument.</text:p><text:list><text:list-item><text:p>Observación documentada</text:p></text:list-item></text:list></office:text></office:body></office:document-content>''')
